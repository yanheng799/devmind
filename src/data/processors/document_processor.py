"""Document processor for parsing and chunking documents."""

import hashlib
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from config import get_settings


logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Parse and chunk documents for vectorization.

    Supports TXT, MD, PDF, and DOCX formats.
    """

    # Magic bytes for file validation
    MAGIC_BYTES = {
        b"%PDF": "pdf",
        b"PK\x03\x04": "docx",
    }

    def __init__(
        self,
        chunk_size: int | None = None,
        chunk_overlap: int | None = None,
    ) -> None:
        """Initialize the document processor.

        Args:
            chunk_size: Target chunk size in characters
            chunk_overlap: Overlap between chunks in characters
        """
        settings = get_settings()
        doc_config = settings.get_document_config()

        self.chunk_size = chunk_size or doc_config["chunk_size"]
        self.chunk_overlap = chunk_overlap or doc_config["chunk_overlap"]
        self.supported_formats = set(doc_config["supported_formats"])
        self.max_file_size = doc_config["max_file_size"]

    def validateFile(self, file_path: str | Path) -> tuple[bool, str | None]:
        """Validate file can be processed.

        Args:
            file_path: Path to file

        Returns:
            Tuple of (is_valid, error_message)
        """
        path = Path(file_path)

        # Check file exists
        if not path.exists():
            return False, f"File not found: {file_path}"

        # Check file is regular file
        if not path.is_file():
            return False, f"Not a regular file: {file_path}"

        # Check file size
        file_size = path.stat().st_size
        if file_size > self.max_file_size:
            return False, (
                f"File too large: {file_size} bytes "
                f"(max {self.max_file_size} bytes)"
            )

        # Check file extension
        ext = path.suffix.lstrip(".").lower()
        if ext not in self.supported_formats:
            return False, (
                f"Unsupported file type: {ext}. "
                f"Supported: {', '.join(sorted(self.supported_formats))}"
            )

        # Validate magic bytes for binary files
        if ext in {"pdf", "docx"}:
            try:
                with open(path, "rb") as f:
                    header = f.read(8)
                    is_valid = False
                    for magic, expected_ext in self.MAGIC_BYTES.items():
                        if header.startswith(magic):
                            if expected_ext == ext:
                                is_valid = True
                            break
                    if not is_valid:
                        return False, f"Invalid file format for .{ext} file"
            except OSError as e:
                return False, f"Failed to read file: {e}"

        return True, None

    def calculateChecksum(self, file_path: str | Path) -> str:
        """Calculate SHA256 checksum of file.

        Args:
            file_path: Path to file

        Returns:
            Hexadecimal checksum string
        """
        sha256 = hashlib.sha256()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(8192), b""):
                sha256.update(chunk)
        return sha256.hexdigest()

    def parseFile(self, file_path: str | Path) -> dict[str, Any]:
        """Parse document and extract metadata and text.

        Args:
            file_path: Path to document

        Returns:
            Dict with keys: text, title, author, metadata
        """
        path = Path(file_path)
        ext = path.suffix.lstrip(".").lower()

        is_valid, error = self.validateFile(path)
        if not is_valid:
            raise ValueError(error)

        if ext == "pdf":
            return self._parse_pdf(path)
        elif ext == "docx":
            return self._parse_docx(path)
        elif ext in {"txt", "md"}:
            return self._parse_text(path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    def _parse_pdf(self, path: Path) -> dict[str, Any]:
        """Parse PDF file.

        Args:
            path: Path to PDF file

        Returns:
            Dict with text and metadata
        """
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError(
                "pypdf is required for PDF processing. "
                "Install with: pip install pypdf"
            )

        reader = PdfReader(path)
        text_parts = []
        metadata: dict[str, Any] = {"page_count": len(reader.pages)}

        # Try to extract metadata
        info = reader.metadata or {}
        title = info.get("/Title", path.stem)
        author = info.get("/Author")
        if author:
            metadata["author"] = author

        # Extract text from all pages
        for page_num, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            except Exception as e:
                logger.warning(f"Failed to extract page {page_num}: {e}")

        text = "\n\n".join(text_parts)

        return {
            "text": text,
            "title": title,
            "author": author,
            "metadata": metadata,
        }

    def _parse_docx(self, path: Path) -> dict[str, Any]:
        """Parse DOCX file.

        Args:
            path: Path to DOCX file

        Returns:
            Dict with text and metadata
        """
        try:
            from docx import Document
        except ImportError:
            raise ImportError(
                "python-docx is required for DOCX processing. "
                "Install with: pip install python-docx"
            )

        doc = Document(path)
        text_parts = []
        metadata: dict[str, Any] = {"paragraph_count": len(doc.paragraphs)}

        # Extract core properties
        core_props = doc.core_properties
        title = core_props.title or path.stem
        author = core_props.author
        if author:
            metadata["author"] = author

        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    text_parts.append(row_text)

        text = "\n\n".join(text_parts)

        return {
            "text": text,
            "title": title,
            "author": author,
            "metadata": metadata,
        }

    def _parse_text(self, path: Path) -> dict[str, Any]:
        """Parse plain text or markdown file.

        Args:
            path: Path to text file

        Returns:
            Dict with text and metadata
        """
        # Try different encodings
        encodings = ["utf-8", "gbk", "gb2312", "big5"]
        text = None
        used_encoding = None

        for encoding in encodings:
            try:
                with open(path, "r", encoding=encoding) as f:
                    text = f.read()
                used_encoding = encoding
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            raise ValueError(
                f"Failed to decode file with any of: {', '.join(encodings)}"
            )

        # Extract title from markdown or use filename
        title = path.stem
        if path.suffix.lower() == ".md":
            # Try to extract first heading
            match = re.match(r"^#\s+(.+)$", text, re.MULTILINE)
            if match:
                title = match.group(1).strip()

        metadata = {"encoding": used_encoding}

        return {
            "text": text,
            "title": title,
            "author": None,
            "metadata": metadata,
        }

    def chunkText(self, text: str) -> list[str]:
        """Split text into overlapping chunks.

        Tries to split at sentence boundaries when possible.

        Args:
            text: Input text

        Returns:
            List of text chunks
        """
        if not text or not text.strip():
            return []

        # Normalize line endings
        text = text.replace("\r\n", "\n").replace("\r", "\n")

        # Split into paragraphs first
        paragraphs = re.split(r"\n\s*\n", text)
        paragraphs = [p.strip() for p in paragraphs if p.strip()]

        if not paragraphs:
            # If no paragraphs, split by lines
            paragraphs = [line.strip() for line in text.split("\n") if line.strip()]

        chunks: list[str] = []
        current_chunk = ""

        for paragraph in paragraphs:
            # If paragraph itself is too long, split it
            if len(paragraph) > self.chunk_size * 2:
                sub_chunks = self._splitLongParagraph(paragraph)
                for sub_chunk in sub_chunks:
                    if len(current_chunk) + len(sub_chunk) + 2 > self.chunk_size:
                        if current_chunk:
                            chunks.append(current_chunk.strip())
                        current_chunk = sub_chunk
                    else:
                        if current_chunk:
                            current_chunk += "\n\n" + sub_chunk
                        else:
                            current_chunk = sub_chunk
                continue

            # Check if adding this paragraph would exceed chunk size
            if len(current_chunk) + len(paragraph) + 2 > self.chunk_size:
                if current_chunk:
                    chunks.append(current_chunk.strip())

                # Handle overlap
                if self.chunk_overlap > 0 and chunks:
                    overlap_text = chunks[-1][-self.chunk_overlap:]
                    current_chunk = overlap_text + "\n\n" + paragraph
                else:
                    current_chunk = paragraph
            else:
                if current_chunk:
                    current_chunk += "\n\n" + paragraph
                else:
                    current_chunk = paragraph

        # Add final chunk
        if current_chunk:
            chunks.append(current_chunk.strip())

        return chunks

    def _splitLongParagraph(self, paragraph: str) -> list[str]:
        """Split a paragraph that exceeds chunk size.

        Tries to split at sentence boundaries.

        Args:
            paragraph: Long paragraph

        Returns:
            List of sub-chunks
        """
        chunks: list[str] = []

        # Try sentence splitting
        sentences = re.split(r"(?<=[.!?。！？])\s+", paragraph)

        current_chunk = ""
        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue

            if len(current_chunk) + len(sentence) + 1 > self.chunk_size:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                current_chunk = sentence
            else:
                if current_chunk:
                    current_chunk += " " + sentence
                else:
                    current_chunk = sentence

        if current_chunk:
            chunks.append(current_chunk.strip())

        # Fallback: if sentences are still too long, force split
        if not chunks or len(chunks[0]) > self.chunk_size * 2:
            chunks = []
            for i in range(0, len(paragraph), self.chunk_size - self.chunk_overlap):
                chunks.append(paragraph[i:i + self.chunk_size])

        return chunks


class MockDocumentProcessor(DocumentProcessor):
    """Mock document processor for testing.

    Returns simple predictable results without actual file parsing.
    """

    def validateFile(self, file_path: str | Path) -> tuple[bool, str | None]:
        """Validate mock file."""
        path = Path(file_path)

        # For testing, accept any file with supported extension
        ext = path.suffix.lstrip(".").lower()
        if ext in self.supported_formats:
            return True, None

        return False, f"Unsupported file type for mock: {ext}"

    def calculateChecksum(self, file_path: str | Path) -> str:
        """Return deterministic mock checksum."""
        path = Path(file_path)
        hash_input = str(path) + "_mock"
        return hashlib.sha256(hash_input.encode()).hexdigest()

    def parseFile(self, file_path: str | Path) -> dict[str, Any]:
        """Return mock parsed content."""
        path = Path(file_path)
        ext = path.suffix.lstrip(".").lower()

        return {
            "text": f"Mock content for {path.name}. " * 100,
            "title": path.stem,
            "author": None,
            "metadata": {
                "page_count": 10 if ext == "pdf" else None,
                "mock": True,
            },
        }


__all__ = [
    "DocumentProcessor",
    "MockDocumentProcessor",
]
